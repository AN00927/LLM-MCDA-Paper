#!/usr/bin/env python
# -*- coding: utf-8 -*-
"""
Deterministic rebuild of the derived scenario files from the three standalone
MASTER workbooks.

Replaces build_consolidated_scenario_workbooks.py, which operated on a
consolidated workbook (ConsolidatedforSimaltaneousediting.xlsx) that no longer
exists. This script reads ONLY:

  masters     HVACScenarios.xlsx, ApplianceScenarios.xlsx, ShowerScenarios.xlsx
  membership  rag_membership.json  (which master rows are RAG; committed, fixed)
  old RAG     HVACRagScenarios.xlsx, ApplianceRAGScenarios.xlsx,
              ShowerRAGScenarios.xlsx  (ONLY to carry their score columns)

and writes:

  TestScenarios.xlsx          195 rows, wide, LLM-facing (banded labels)
  HVACRagScenarios.xlsx       35 scenarios x 3 rows, long, carried scores
  ApplianceRAGScenarios.xlsx  35 scenarios x 3 rows
  ShowerRAGScenarios.xlsx     20 scenarios x 3 rows

Partition. RAG membership is NOT re-derived by matching the old RAG sheets each
time (that is what made the old script fragile: a fixed question could orphan
its RAG scenario). It is read from rag_membership.json, which records each RAG
scenario's id and a master-row signature: the scenario parameters plus the
alternative SET, with the question text and the engineering-only values
(hvac_age, r_value, seer, tank_size, water_heater_temp) deliberately left out.
So a grammar fix to a question, a cell-type fix, or an hvac_age correction never
changes which scenarios are RAG. A fix to a signature field (e.g. a location
typo in a RAG row) fails loudly instead of silently moving the scenario.
Create the file once with --init-membership (derives it from the current RAG
sheets); after that it is the source of truth for the split.

Carried scores. The five score columns, rank and raw_* are copied verbatim from
the existing RAG sheet row with the same scenario_id and alternative (after the
scenario's signature is checked against the membership entry). They are
refreshed afterwards by sync_rag_ground_truth_scores.py, exactly as before.

Formats. TestScenarios is written with explicit Text / Number formats ("@",
"0", "0.######"), as the old export did. The RAG sheets are written through the
same pandas writer sync_rag_ground_truth_scores.py uses (General formats),
because the sync always rewrites them that way right after a rebuild; emitting
that format here makes rebuild -> sync a no-op on everything but score values.

Safety.
  - Default output is a fresh scratch directory; --out-dir picks another one.
  - Only --in-place writes into the masters directory (Scenario Files/), and
    only after the audit passes, backing up each file to Backups/ first.
  - The masters are never written.

Refresh workflow (CLAUDE.md): edit masters -> this script -> GT calculators ->
sync_rag_ground_truth_scores.py -> build_rag_index.py (always last).

Usage:
    python "Scenario Files/build_scenario_workbooks.py"                  # -> scratch dir
    python "Scenario Files/build_scenario_workbooks.py" --out-dir DIR
    python "Scenario Files/build_scenario_workbooks.py" --masters-dir DIR --out-dir DIR2
    python "Scenario Files/build_scenario_workbooks.py" --in-place       # live refresh
    python "Scenario Files/build_scenario_workbooks.py" --init-membership  # one-time
"""

import argparse
import importlib.util
import json
import os
import re
import shutil
import sys
import tempfile
import unicodedata
from collections import Counter, defaultdict
from datetime import datetime, timezone
from pathlib import Path

import openpyxl
import pandas as pd

HERE = os.path.dirname(os.path.abspath(__file__))
PROJECT_ROOT = os.path.dirname(HERE)
if PROJECT_ROOT not in sys.path:
    sys.path.insert(0, PROJECT_ROOT)
# Single source of truth for the banded labels (also used by the RAG embedding).
from sentinel_utils import (appliance_age_to_band_label, gpm_to_flow_rate_label,
                            house_age_to_band_label, is_sentinel)

DEFAULT_MEMBERSHIP = os.path.join(HERE, "rag_membership.json")
SYNC_SCRIPT = os.path.join(PROJECT_ROOT, "Miscellaneous Scripts", "core-automation",
                           "sync_rag_ground_truth_scores.py")

MASTER_FILES = {"HVAC": "HVACScenarios.xlsx", "Appliance": "ApplianceScenarios.xlsx",
                "Shower": "ShowerScenarios.xlsx"}
RAG_FILES = {"HVAC": "HVACRagScenarios.xlsx", "Appliance": "ApplianceRAGScenarios.xlsx",
             "Shower": "ShowerRAGScenarios.xlsx"}
TEST_FILE = "TestScenarios.xlsx"
TYPES = ["HVAC", "Appliance", "Shower"]

# Fixed benchmark split: (test scenarios, RAG scenarios) per decision type.
EXPECTED_COUNTS = {"HVAC": (70, 35), "Appliance": (65, 35), "Shower": (60, 20)}

# ----------------------------------------------------------------------------
# normalisation helpers (same rules as the old consolidated rebuild)
# ----------------------------------------------------------------------------
REPLACEMENT = chr(0xFFFD)                           # legacy corruption marker
EM, EN, NBSP = chr(0x2014), chr(0x2013), chr(0xA0)  # em-dash, en-dash, nbsp


def clean_text(v):
    """Fold nbsp, em-dash -> ', ', digit-flanked en-dash -> '-', other en-dash
    -> ', ', collapse whitespace, trim. Empty -> None."""
    if v is None:
        return None
    s = str(v).replace(REPLACEMENT, EM).replace(NBSP, " ")
    s = re.sub(r"\s*" + EM + r"\s*", ", ", s)
    s = re.sub(r"(?<=\d)" + EN + r"(?=\d)", "-", s)
    s = re.sub(r"\s*" + EN + r"\s*", ", ", s)
    s = re.sub(r"\s+", " ", s).strip()
    return s if s != "" else None


def has_corruption(s):
    """True if a string holds U+FFFD or a control/format char (tab/newline ok)."""
    if not isinstance(s, str):
        return False
    if REPLACEMENT in s:
        return True
    return any(unicodedata.category(ch).startswith("C") for ch in s if ch not in "\t\n")


def to_num(v):
    """Numeric or numeric-string -> int when integral, else float."""
    if v is None:
        return None
    if isinstance(v, (int, float)) and not isinstance(v, bool):
        f = float(v)
    else:
        f = float(clean_text(v))
    return int(f) if f.is_integer() else f


_TIME_FULL = re.compile(r"^(\d{1,2}):(\d{2})\s*([ap])\.?m\.?$", re.I)
_TIME_HOUR = re.compile(r"^(\d{1,2})\s*([ap])\.?m\.?$", re.I)


def to_clock(v):
    """Clock value -> 'H:MM AM/PM' text. Handles Excel day-fraction serials and
    '6pm' / '8am' shorthand. Unparseable text is returned cleaned, unchanged."""
    if v is None:
        return None
    if isinstance(v, (int, float)) and not isinstance(v, bool):
        total = int(round((float(v) % 1.0) * 24 * 60))
        h, m = (total // 60) % 24, total % 60
        return f"{h % 12 or 12}:{m:02d} {'AM' if h < 12 else 'PM'}"
    s = clean_text(v)
    if s is None:
        return None
    mo = _TIME_FULL.match(s)
    if mo:
        return f"{int(mo.group(1))}:{int(mo.group(2)):02d} {mo.group(3).upper()}M"
    mo = _TIME_HOUR.match(s)
    if mo:
        return f"{int(mo.group(1))}:00 {mo.group(2).upper()}M"
    return s


def alt_norm(v, kind):
    """Alternative as stored in the derived sheets: 'clock' -> 'H:MM AM/PM';
    'int' -> integer string ('80'), or cleaned text for 'Off'."""
    if kind == "clock":
        return to_clock(v)
    s = clean_text(v) if isinstance(v, str) else v
    try:
        return str(int(round(float(s))))
    except (TypeError, ValueError):
        return clean_text(v)


def canon(v):
    """Comparison form: numbers -> float rounded to 6 dp, text -> cleaned str."""
    if v is None:
        return None
    if isinstance(v, (int, float)) and not isinstance(v, bool):
        return round(float(v), 6)
    s = clean_text(v)
    if s is None:
        return None
    try:
        return round(float(s), 6)
    except ValueError:
        return s


APPLIANCE_MAP = {"washer": "washing_machine"}


def canon_appliance(v):
    s = clean_text(v)
    s = s.lower() if s else s
    return APPLIANCE_MAP.get(s, s)


ALT_KIND = {"HVAC": "int", "Appliance": "clock", "Shower": "int"}
KNOWN_NONINT_ALTS = {"off"}   # HVAC "turn the system off" setpoint

# ----------------------------------------------------------------------------
# output schemas: (column, cell kind). Kinds drive typing + number format.
# ----------------------------------------------------------------------------
FMT = {"text": "@", "int": "0", "num": "0.######"}

TEST_COLS = [
    ("decision_type", "text"), ("question", "text"), ("location", "text"),
    ("square_footage", "num"), ("insulation", "text"), ("household_size", "int"),
    ("utility_budget", "num"), ("housing_type", "text"), ("outdoor_temp", "num"),
    ("house_age", "text"), ("appliance_age", "text"), ("flow_rate", "text"),
    ("alternative_1", "text"), ("alternative_2", "text"), ("alternative_3", "text"),
]
SCORE_COLS = ["energy_cost_score", "environmental_score", "comfort_score",
              "practicality_score", "mavt_score", "rank"]
RAW_COLS = {"HVAC": ["raw_kwh", "raw_cost", "raw_emissions"],
            "Appliance": ["raw_cost", "raw_emissions"],
            "Shower": ["raw_kwh", "raw_cost", "raw_water_gallons"]}
RAG_COLS = {
    "HVAC": [
        ("scenario_id", "int"), ("question", "text"), ("location", "text"),
        ("square_footage", "num"), ("insulation", "text"), ("household_size", "int"),
        ("utility_budget", "num"), ("housing_type", "text"), ("outdoor_temp", "num"),
        ("house_age", "num"), ("r_value", "num"), ("seer", "num"), ("hvac_age", "num"),
        ("alternative", "text"),
    ],
    "Appliance": [
        ("scenario_id", "int"), ("question", "text"), ("location", "text"),
        ("utility_budget", "num"), ("appliance", "text"), ("appliance_age", "text"),
        ("housing_type", "text"), ("household_size", "int"), ("kwh_per_cycle", "num"),
        ("alternative", "text"),
    ],
    "Shower": [
        ("scenario_id", "int"), ("question", "text"), ("location", "text"),
        ("household_size", "int"), ("gpm", "num"), ("flow_rate", "text"),
        ("tank_size", "num"), ("water_heater_temp", "num"), ("utility_budget", "num"),
        ("housing_type", "text"), ("outdoor_temp", "num"), ("alternative", "text"),
        ("duration_min", "int"),
    ],
}
for _t in TYPES:  # score / raw columns are carried verbatim (kind "carry")
    RAG_COLS[_t] = RAG_COLS[_t] + [(c, "carry") for c in SCORE_COLS + RAW_COLS[_t]]

# Membership signature fields: the scenario parameters that identify a master
# row. Question text and engineering-only values are excluded on purpose.
SIG_FIELDS = {
    "HVAC": ["location", "square_footage", "insulation", "household_size",
             "utility_budget", "housing_type", "outdoor_temp", "house_age"],
    "Appliance": ["location", "utility_budget", "appliance", "housing_type",
                  "household_size", "kwh_per_cycle", "appliance_age"],
    "Shower": ["location", "household_size", "gpm", "utility_budget",
               "housing_type", "outdoor_temp"],
}


def typed(kind, val):
    """Python value a cell of this kind stores (None for blank)."""
    if val is None or (isinstance(val, str) and val.strip() == ""):
        return None
    if kind == "text":
        return str(val)
    if kind == "int":
        return int(round(float(val)))
    if kind == "num":
        f = float(val)
        return int(f) if f.is_integer() else f
    return val  # "carry": verbatim


# ----------------------------------------------------------------------------
# reading
# ----------------------------------------------------------------------------
def read_rows(path):
    """First worksheet -> list of (0-based data index, {header: value}).
    Fully blank rows are skipped (their index is still consumed)."""
    ws = openpyxl.load_workbook(path).worksheets[0]
    hdr = [c.value for c in ws[1]]
    out = []
    for i, row in enumerate(ws.iter_rows(min_row=2, values_only=True)):
        if all(v is None or (isinstance(v, str) and v.strip() == "") for v in row):
            continue
        out.append((i, dict(zip(hdr, row))))
    return out


class AuditError(Exception):
    pass


def load_masters(masters_dir, notes):
    """Normalise the three masters in memory. Returns {type: [record]}; a
    record keeps the master's 0-based data index as 'row'. Never writes."""
    masters = {}
    for t in TYPES:
        path = os.path.join(masters_dir, MASTER_FILES[t])
        recs = []
        changed = Counter()   # how many master cells normalisation altered
        for i, d in read_rows(path):
            def txt(col):
                v = clean_text(d[col])
                if v != d[col]:
                    changed[col] += 1
                return v
            alts = []
            for n in (1, 2, 3):
                raw = d[f"alternative_{n}"]
                a = alt_norm(raw, ALT_KIND[t])
                if str(raw) != str(a):
                    changed[f"alternative_{n}"] += 1
                alts.append(a)
            rec = {"row": i, "question": txt("question"), "location": txt("location"),
                   "housing_type": txt("housing_type"),
                   "household_size": int(to_num(d["household_size"])),
                   "utility_budget": to_num(d["utility_budget"]), "alt": alts}
            if t == "HVAC":
                rec.update({
                    "square_footage": to_num(d["square_footage"]),
                    "insulation": txt("insulation"),
                    "outdoor_temp": to_num(d["outdoor_temp"]),
                    # house_age kept as cleaned text like the old rebuild
                    # (banded for Test, numeric-typed for RAG)
                    "house_age": clean_text(str(d["house_age"])),
                    "r_value": to_num(d["r_value"]), "seer": to_num(d["seer"]),
                    "hvac_age": to_num(d["hvac_age"]),
                })
            elif t == "Appliance":
                appv = canon_appliance(d["appliance"])
                if appv != d["appliance"]:
                    changed["appliance"] += 1
                bt = to_clock(d["baseline_time"])
                if str(bt) != str(d["baseline_time"]):
                    changed["baseline_time"] += 1
                rec.update({
                    "appliance": appv, "baseline_time": bt,
                    "kwh_per_cycle": to_num(d["kwh_per_cycle"]),
                    "appliance_age": clean_text(str(d["appliance_age"])),
                })
            else:
                rec.update({
                    "gpm": to_num(d["gpm"]), "outdoor_temp": to_num(d["outdoor_temp"]),
                    "tank_size": to_num(d["tank_size"]),
                    "water_heater_temp": to_num(d["water_heater_temp"]),
                })
            recs.append(rec)
        if changed:
            notes.append(f"{MASTER_FILES[t]}: in-memory normalisation changed "
                         f"{dict(changed)} cell(s) (masters themselves not written)")
        masters[t] = recs
    return masters


def signature(t, rec_or_row, alts):
    """Membership signature as a JSON-friendly dict (canonical values)."""
    sig = {}
    for k in SIG_FIELDS[t]:
        v = rec_or_row[k]
        sig[k] = canon_appliance(v) if k == "appliance" else canon(v)
    sig["alternatives"] = sorted((alt_norm(a, ALT_KIND[t]) for a in alts), key=str)
    return sig


def sig_key(sig):
    return json.dumps(sig, sort_keys=True)


def read_rag_groups(rag_dir, t):
    """Existing RAG sheet -> {scenario_id: [row dicts in file order]}."""
    groups = defaultdict(list)
    for _, d in read_rows(os.path.join(rag_dir, RAG_FILES[t])):
        groups[d["scenario_id"]].append(d)
    return groups


# ----------------------------------------------------------------------------
# membership
# ----------------------------------------------------------------------------
def init_membership(masters, rag_dir, path):
    """One-time: derive membership from the CURRENT RAG sheets (each RAG group
    must match exactly one master row by signature) and write it to *path*."""
    out = {"description": ("RAG/Test partition for build_scenario_workbooks.py. Each "
                           "entry is one RAG scenario: its scenario_id and the master-row "
                           "signature (parameters + alternative set; question text and "
                           "engineering-only values excluded). master_row (0-based data "
                           "index) and question are informational."),
           "created_utc": datetime.now(timezone.utc).strftime("%Y-%m-%dT%H:%M:%SZ"),
           "created_from": "current RAG sheets in " + os.path.relpath(rag_dir, PROJECT_ROOT),
           "types": {}}
    for t in TYPES:
        idx = defaultdict(list)
        for m in masters[t]:
            idx[sig_key(signature(t, m, m["alt"]))].append(m)
        entries = []
        for sid, grp in sorted(read_rag_groups(rag_dir, t).items()):
            sig = signature(t, grp[0], [g["alternative"] for g in grp])
            hits = idx.get(sig_key(sig), [])
            if len(hits) != 1:
                raise AuditError(f"init-membership: {t} RAG scenario_id {sid} matches "
                                 f"{len(hits)} master rows (need exactly 1)")
            entries.append({"scenario_id": int(sid), "master_row": hits[0]["row"],
                            "question": hits[0]["question"], "signature": sig})
        out["types"][t] = entries
    with open(path, "w", encoding="utf-8") as f:
        json.dump(out, f, indent=1, ensure_ascii=True)
        f.write("\n")
    return out


def resolve_membership(masters, membership, notes):
    """Map each membership entry to exactly one master row.
    Returns {type: [(scenario_id, master_record), ...] sorted by scenario_id}."""
    resolved = {}
    for t in TYPES:
        idx = defaultdict(list)
        for m in masters[t]:
            idx[sig_key(signature(t, m, m["alt"]))].append(m)
        dup = [k for k, v in idx.items() if len(v) > 1]
        if dup:
            raise AuditError(f"{t}: {len(dup)} master signature(s) shared by >1 row, e.g. "
                             f"rows {[m['row'] for m in idx[dup[0]]]}")
        pairs, claimed = [], set()
        for e in membership["types"][t]:
            hits = idx.get(sig_key(e["signature"]), [])
            if len(hits) != 1:
                raise AuditError(
                    f"{t} RAG scenario_id {e['scenario_id']}: membership signature matches "
                    f"{len(hits)} master rows (recorded master_row {e['master_row']}). A "
                    f"signature field of that master row was edited; update "
                    f"rag_membership.json deliberately if the edit is intended.")
            m = hits[0]
            if m["row"] in claimed:
                raise AuditError(f"{t}: master row {m['row']} claimed twice by membership")
            claimed.add(m["row"])
            if m["row"] != e["master_row"]:
                notes.append(f"{t} RAG id {e['scenario_id']}: master row moved "
                             f"{e['master_row']} -> {m['row']} (matched by signature)")
            pairs.append((int(e["scenario_id"]), m))
        pairs.sort(key=lambda p: p[0])
        resolved[t] = pairs
    return resolved


# ----------------------------------------------------------------------------
# building the derived rows
# ----------------------------------------------------------------------------
def build_rag_rows(t, pairs, old_groups):
    """Long RAG rows (3 per scenario, master alternative order) with the score
    columns carried from the existing RAG row of the same id + alternative."""
    rows = []
    for sid, m in pairs:
        old = old_groups.get(sid)
        if not old:
            raise AuditError(f"{t} RAG id {sid}: no rows in the existing RAG sheet to "
                             f"carry scores from")
        old_sig = signature(t, old[0], [g["alternative"] for g in old])
        if sig_key(old_sig) != sig_key(signature(t, m, m["alt"])):
            raise AuditError(f"{t} RAG id {sid}: existing RAG sheet row does not match the "
                             f"master row it belongs to; cannot carry scores safely")
        bundles = {alt_norm(g["alternative"], ALT_KIND[t]): g for g in old}
        for a in m["alt"]:
            b = bundles[a]
            row = {"scenario_id": sid, "question": m["question"], "location": m["location"],
                   "household_size": m["household_size"], "utility_budget": m["utility_budget"],
                   "housing_type": m["housing_type"], "alternative": a}
            if t == "HVAC":
                row.update({k: m[k] for k in ("square_footage", "insulation", "outdoor_temp",
                                              "house_age", "r_value", "seer", "hvac_age")})
            elif t == "Appliance":
                row.update({k: m[k] for k in ("appliance", "appliance_age", "kwh_per_cycle")})
            else:
                row.update({k: m[k] for k in ("gpm", "tank_size", "water_heater_temp",
                                              "outdoor_temp")})
                row["flow_rate"] = gpm_to_flow_rate_label(m["gpm"])
                row["duration_min"] = int(a)
            for c in SCORE_COLS + RAW_COLS[t]:
                row[c] = b[c]
            rows.append({name: typed(kind, row.get(name)) for name, kind in RAG_COLS[t]})
    return rows


def build_test_rows(masters, resolved):
    """Wide Test rows: every master row NOT in RAG, masters in HVAC / Appliance /
    Shower order, master row order within type. Returns (rows, source rows)."""
    rows, sources = [], []
    for t in TYPES:
        rag_rows = {m["row"] for _, m in resolved[t]}
        for m in masters[t]:
            if m["row"] in rag_rows:
                continue
            row = {"decision_type": t, "question": m["question"], "location": m["location"],
                   "household_size": m["household_size"], "utility_budget": m["utility_budget"],
                   "housing_type": m["housing_type"],
                   "alternative_1": m["alt"][0], "alternative_2": m["alt"][1],
                   "alternative_3": m["alt"][2]}
            if t == "HVAC":
                row.update({"square_footage": m["square_footage"], "insulation": m["insulation"],
                            "outdoor_temp": m["outdoor_temp"],
                            "house_age": house_age_to_band_label(m["house_age"])})
            elif t == "Appliance":
                row["appliance_age"] = appliance_age_to_band_label(m["appliance_age"])
            else:
                row.update({"outdoor_temp": m["outdoor_temp"],
                            "flow_rate": gpm_to_flow_rate_label(m["gpm"])})
            rows.append(row)
            sources.append((t, m["row"]))
    return rows, sources


# ----------------------------------------------------------------------------
# writing
# ----------------------------------------------------------------------------
def write_test(path, rows):
    """TestScenarios: explicit Text/Number formats on every cell, blanks too."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    for ci, (name, _) in enumerate(TEST_COLS, start=1):
        c = ws.cell(1, ci)
        c.value, c.number_format = name, "@"
    for ri, row in enumerate(rows, start=2):
        for ci, (name, kind) in enumerate(TEST_COLS, start=1):
            c = ws.cell(ri, ci)
            c.value = typed(kind, row.get(name))
            c.number_format = FMT[kind]
    wb.save(path)


def load_sync_writer():
    """write_xlsx from sync_rag_ground_truth_scores.py (the RAG sheets' writer)."""
    spec = importlib.util.spec_from_file_location("sync_rag_ground_truth_scores", SYNC_SCRIPT)
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    return mod.write_xlsx


def write_rag(path, t, rows, sync_write):
    df = pd.DataFrame(rows, columns=[c for c, _ in RAG_COLS[t]], dtype=object)
    sync_write(df, Path(path))


# ----------------------------------------------------------------------------
# audit (runs on the files as written, read back from disk)
# ----------------------------------------------------------------------------
def audit(out_dir, masters, resolved, test_sources):
    """Returns list of (id, level, message). Any FAIL aborts the run."""
    res = []

    def A(cid, ok, msg_ok, msg_bad):
        res.append((cid, "PASS" if ok else "FAIL", msg_ok if ok else msg_bad))

    ws = openpyxl.load_workbook(os.path.join(out_dir, TEST_FILE)).worksheets[0]
    test_hdr = [c.value for c in ws[1]]
    test = [dict(zip(test_hdr, r)) for r in ws.iter_rows(min_row=2, values_only=True)]
    test_cells = list(ws.iter_rows(min_row=2))
    rag = {t: read_rag_groups(out_dir, t) for t in TYPES}
    rag_hdr = {t: [c.value for c in openpyxl.load_workbook(
        os.path.join(out_dir, RAG_FILES[t])).worksheets[0][1]] for t in TYPES}

    # A01 headers
    bad = [TEST_FILE] if test_hdr != [c for c, _ in TEST_COLS] else []
    bad += [RAG_FILES[t] for t in TYPES if rag_hdr[t] != [c for c, _ in RAG_COLS[t]]]
    A("A01", not bad, "canonical headers on all four sheets", f"header drift in {bad}")

    # A02 counts: fixed split and partition complete
    lines, ok = [], True
    for t in TYPES:
        n_test = sum(1 for r in test if r["decision_type"] == t)
        n_rag = len(rag[t])
        n_m = len(masters[t])
        exp = EXPECTED_COUNTS[t]
        good = (n_test, n_rag) == exp and n_test + n_rag == n_m
        ok &= good
        lines.append(f"{t} test={n_test} rag={n_rag} master={n_m} (expect {exp[0]}/{exp[1]})")
    A("A02", ok, "; ".join(lines), "count mismatch: " + "; ".join(lines))

    # A03 every Test row traces to exactly one master row (transform rules applied)
    def test_match(r, m, t):
        want = {"question": m["question"], "location": m["location"],
                "household_size": m["household_size"], "utility_budget": m["utility_budget"],
                "housing_type": m["housing_type"], "alternative_1": m["alt"][0],
                "alternative_2": m["alt"][1], "alternative_3": m["alt"][2],
                "square_footage": None, "insulation": None, "outdoor_temp": None,
                "house_age": None, "appliance_age": None, "flow_rate": None}
        if t == "HVAC":
            want.update(square_footage=m["square_footage"], insulation=m["insulation"],
                        outdoor_temp=m["outdoor_temp"],
                        house_age=house_age_to_band_label(m["house_age"]))
        elif t == "Appliance":
            want["appliance_age"] = appliance_age_to_band_label(m["appliance_age"])
        else:
            want.update(outdoor_temp=m["outdoor_temp"], flow_rate=gpm_to_flow_rate_label(m["gpm"]))
        return all(canon(r[k]) == canon(v) for k, v in want.items())

    test_claims, bad = defaultdict(set), []
    for i, r in enumerate(test):
        t = r["decision_type"]
        hits = [m["row"] for m in masters.get(t, []) if test_match(r, m, t)]
        # must be exactly one master row, and the one this Test row was built from
        if len(hits) != 1 or (t, hits[0]) != test_sources[i]:
            bad.append((i + 2, t, hits))
        else:
            test_claims[t].add(hits[0])
    A("A03", not bad and len(test) == len(test_sources),
      "every Test row traces to exactly one master row (the one it was built from)",
      f"{len(bad)} Test row(s) match != 1 master row, e.g. {bad[:5]}")

    # A04 every RAG scenario traces to exactly one master row (raw values, question,
    #     engineering values, alternatives in master order)
    def rag_match(grp, m, t):
        keys = ["question", "location", "household_size", "utility_budget", "housing_type"]
        keys += {"HVAC": ["square_footage", "insulation", "outdoor_temp", "house_age",
                          "r_value", "seer", "hvac_age"],
                 "Appliance": ["appliance", "appliance_age", "kwh_per_cycle"],
                 "Shower": ["gpm", "tank_size", "water_heater_temp", "outdoor_temp"]}[t]
        if any(canon(g[k]) != canon(m[k]) for g in grp for k in keys):
            return False
        if t == "Shower" and any(g["flow_rate"] != gpm_to_flow_rate_label(m["gpm"]) for g in grp):
            return False
        return [str(g["alternative"]) for g in grp] == m["alt"]

    rag_claims, bad = defaultdict(set), []
    for t in TYPES:
        for sid, grp in rag[t].items():
            hits = [m["row"] for m in masters[t] if rag_match(grp, m, t)]
            if len(hits) != 1:
                bad.append((t, sid, len(hits)))
            else:
                rag_claims[t].add(hits[0])
    A("A04", not bad, "every RAG scenario traces to exactly one master row "
                      "(alternatives in master order)",
      f"{len(bad)} RAG scenario(s) match != 1 master row, e.g. {bad[:5]}")

    # A05 disjoint and complete; RAG rows are exactly the membership rows
    lines, ok = [], True
    for t in TYPES:
        all_rows = {m["row"] for m in masters[t]}
        both = test_claims[t] & rag_claims[t]
        missing = all_rows - test_claims[t] - rag_claims[t]
        if both or missing:
            ok = False
            lines.append(f"{t}: in both {sorted(both)[:5]}, in neither {sorted(missing)[:5]}")
        if rag_claims[t] != {m["row"] for _, m in resolved[t]}:
            ok = False
            lines.append(f"{t}: RAG rows differ from membership")
    A("A05", ok, "Test and RAG disjoint and together cover every master row exactly once",
      "partition broken: " + "; ".join(lines))

    # A06 RAG structure: ids 1..N, 3 consecutive rows each
    bad = []
    for t in TYPES:
        ids = [d["scenario_id"] for _, d in read_rows(os.path.join(out_dir, RAG_FILES[t]))]
        uniq = list(dict.fromkeys(ids))
        if uniq != list(range(1, len(uniq) + 1)):
            bad.append(f"{t}: ids not 1..N in order")
        if any(ids.count(s) != 3 for s in uniq):
            bad.append(f"{t}: a scenario without exactly 3 rows")
        if ids != [s for s in uniq for _ in range(3)]:
            bad.append(f"{t}: rows of a scenario not consecutive")
    A("A06", not bad, "RAG ids 1..N per type, 3 consecutive rows each", "; ".join(bad))

    # A07 carried scores numeric, no sentinel, rank consistent with mavt
    bad = []
    for t in TYPES:
        for sid, grp in rag[t].items():
            for g in grp:
                for c in SCORE_COLS + RAW_COLS[t]:
                    v = g[c]
                    if not isinstance(v, (int, float)) or isinstance(v, bool) or is_sentinel(v):
                        bad.append((t, sid, c, v))
            ranks = sorted(g["rank"] for g in grp)
            order = sorted(grp, key=lambda g: g["rank"])
            mv = [float(g["mavt_score"]) for g in order]
            if ranks != [1, 2, 3] or mv != sorted(mv, reverse=True):
                bad.append((t, sid, "rank", ranks))
    A("A07", not bad, "RAG scores/raw populated, numeric, no sentinel; rank consistent with mavt",
      f"{len(bad)} bad score/rank cell(s), e.g. {bad[:5]}")

    # A08 Test cell formats and types
    bad = 0
    for cells in test_cells:
        for c, (name, kind) in zip(cells, TEST_COLS):
            if c.number_format != FMT[kind]:
                bad += 1
            v = c.value
            if v is not None and kind == "text" and not isinstance(v, str):
                bad += 1
            if v is not None and kind != "text" and (not isinstance(v, (int, float))
                                                       or isinstance(v, bool)):
                bad += 1
    A("A08", not bad, "TestScenarios: every cell Text/Number per schema, none General",
      f"TestScenarios: {bad} format/type mismatch(es)")

    # A09 no corruption; enums valid; appliance times are H:MM AM/PM text
    INS = {"Good", "Medium", "Poor"}
    HT = {"Apartment", "Condo", "Rowhouse", "Single-family", "Townhouse"}
    APP = {"dishwasher", "dryer", "washing_machine"}
    clock = re.compile(r"^\d{1,2}:\d{2} [AP]M$")
    bad = []
    everything = test + [g for t in TYPES for grp in rag[t].values() for g in grp]
    for r in everything:
        for k, v in r.items():
            if has_corruption(v):
                bad.append(("corrupt", k, v))
        if r.get("insulation") is not None and r["insulation"] not in INS:
            bad.append(("insulation", r["insulation"]))
        if r.get("housing_type") not in HT:
            bad.append(("housing_type", r.get("housing_type")))
        if "appliance" in r and r["appliance"] not in APP:
            bad.append(("appliance", r["appliance"]))
    for r in test:
        if r["decision_type"] == "Appliance":
            bad += [("time", r[k]) for k in ("alternative_1", "alternative_2", "alternative_3")
                    if not clock.match(str(r[k]))]
    for grp in rag["Appliance"].values():
        bad += [("time", g["alternative"]) for g in grp if not clock.match(str(g["alternative"]))]
    for m in masters["HVAC"]:
        bad += [("hvac_alt", a) for a in m["alt"]
                if not re.match(r"^-?\d+$", str(a)) and str(a).lower() not in KNOWN_NONINT_ALTS]
    A("A09", not bad, "no corrupt text; enums valid; appliance times H:MM AM/PM; "
                      "HVAC alternatives integer or 'Off'",
      f"{len(bad)} problem(s), e.g. {ascii(bad[:6])}")
    return res


# ----------------------------------------------------------------------------
# main
# ----------------------------------------------------------------------------
def main(argv=None):
    ap = argparse.ArgumentParser(description=__doc__.split("\n\n")[0],
                                 formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--masters-dir", default=HERE,
                    help="directory holding the three master workbooks (default: Scenario Files/)")
    ap.add_argument("--rag-scores-dir", default=None,
                    help="directory holding the existing RAG sheets whose score columns are "
                         "carried (default: --masters-dir)")
    ap.add_argument("--membership", default=DEFAULT_MEMBERSHIP,
                    help="RAG membership JSON (default: Scenario Files/rag_membership.json)")
    ap.add_argument("--out-dir", default=None,
                    help="where to write the four derived files (default: a new scratch dir)")
    ap.add_argument("--in-place", action="store_true",
                    help="write into --masters-dir after the audit passes, backing up each "
                         "file to <masters-dir>/Backups/ first")
    ap.add_argument("--init-membership", action="store_true",
                    help="derive the membership file from the current RAG sheets and exit "
                         "(refuses to overwrite an existing file)")
    args = ap.parse_args(argv)

    masters_dir = os.path.abspath(args.masters_dir)
    rag_dir = os.path.abspath(args.rag_scores_dir or masters_dir)
    if args.in_place and args.out_dir:
        ap.error("--in-place and --out-dir are mutually exclusive")

    notes = []
    masters = load_masters(masters_dir, notes)
    print(f"[masters] {masters_dir}: " +
          ", ".join(f"{t}={len(masters[t])}" for t in TYPES))

    if args.init_membership:
        if os.path.exists(args.membership):
            print(f"[init-membership] REFUSED: {args.membership} already exists")
            return 2
        mem = init_membership(masters, rag_dir, args.membership)
        print(f"[init-membership] wrote {args.membership}: " +
              ", ".join(f"{t}={len(mem['types'][t])}" for t in TYPES))
        return 0

    with open(args.membership, "r", encoding="utf-8") as f:
        membership = json.load(f)

    try:
        resolved = resolve_membership(masters, membership, notes)
        rag_rows = {t: build_rag_rows(t, resolved[t], read_rag_groups(rag_dir, t))
                    for t in TYPES}
        test_rows, test_sources = build_test_rows(masters, resolved)
    except AuditError as exc:
        print(f"[FAIL] {exc}")
        print("==================> REBUILD ABORTED, nothing written <==================")
        return 1

    # Stage: in-place writes go to a temp dir first and are copied only on PASS.
    ts = datetime.now(timezone.utc).strftime("%Y%m%dT%H%M%SZ")
    if args.in_place:
        stage = tempfile.mkdtemp(prefix="scenario_rebuild_stage_")
    else:
        stage = os.path.abspath(args.out_dir or os.path.join(
            tempfile.gettempdir(), f"scenario_rebuild_{ts}"))
        os.makedirs(stage, exist_ok=True)

    sync_write = load_sync_writer()
    write_test(os.path.join(stage, TEST_FILE), test_rows)
    for t in TYPES:
        write_rag(os.path.join(stage, RAG_FILES[t]), t, rag_rows[t], sync_write)

    results = audit(stage, masters, resolved, test_sources)
    fails = [r for r in results if r[1] == "FAIL"]

    print("\n==================== AUDIT ====================")
    for cid, lvl, msg in results:
        print(f"  [{lvl:4}] {cid} {msg}")
    for n in notes:
        print(f"  [INFO] {n}")

    if not args.in_place:
        with open(os.path.join(stage, "rebuild_audit.txt"), "w", encoding="utf-8") as f:
            f.write(f"build_scenario_workbooks.py run at {ts} (UTC)\n")
            f.write(f"masters: {masters_dir}\nscores carried from: {rag_dir}\n"
                    f"membership: {os.path.abspath(args.membership)}\n\n")
            for cid, lvl, msg in results:
                f.write(f"[{lvl}] {cid} {msg}\n")
            for n in notes:
                f.write(f"[INFO] {n}\n")

    if fails:
        print(f"\n==================> AUDIT FAILED: {len(fails)} hard violation(s) "
              f"<==================")
        if args.in_place:
            print("[in-place] nothing copied into " + masters_dir)
        else:
            print(f"[out] failing output left in {stage} for inspection")
        return 1

    if args.in_place:
        backup_dir = os.path.join(masters_dir, "Backups")
        os.makedirs(backup_dir, exist_ok=True)
        for fname in [TEST_FILE] + [RAG_FILES[t] for t in TYPES]:
            dest = os.path.join(masters_dir, fname)
            if os.path.exists(dest):
                bak = os.path.join(backup_dir,
                                   f"{os.path.splitext(fname)[0]}_BACKUP_{ts}.xlsx")
                shutil.copy2(dest, bak)
                if not os.path.exists(bak) or os.path.getsize(bak) == 0:
                    print(f"[FAIL] backup of {fname} failed; stopping before any write")
                    return 1
            shutil.copy2(os.path.join(stage, fname), dest)
            print(f"[in-place] {fname} written (backup in Backups/)")
        shutil.rmtree(stage, ignore_errors=True)
        print("Next: GT calculators -> sync_rag_ground_truth_scores.py -> build_rag_index.py")
    else:
        print(f"\n[out] {stage}")
    print("\n==================> AUDIT PASSED <==================")
    return 0


if __name__ == "__main__":
    sys.exit(main())
